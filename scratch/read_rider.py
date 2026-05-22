import docx
import os

docx_path = "Rider sonido CIEDI.docx"
output_path = "scratch/rider_docx_text.txt"

doc = docx.Document(docx_path)

with open(output_path, "w", encoding="utf-8") as f:
    f.write("=== PARAGRAPHS ===\n")
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            f.write(f"[{i}] {para.text}\n")
            
    f.write("\n=== TABLES ===\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"\nTable {t_idx}:\n")
        for row in table.rows:
            row_text = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            f.write(" | ".join(row_text) + "\n")

print("Done extracting Rider docx.")
