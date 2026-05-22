import docx
import json

doc = docx.Document("CEREMONIA DE PRIMARIA/Minuto a Minuto - Primaria Ohana.docx")

print("--- PARAGRAPHS ---")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"{i}: {p.text}")

print("\n--- TABLES ---")
for t_idx, table in enumerate(doc.tables):
    print(f"Table {t_idx}:")
    for r_idx, row in enumerate(table.rows):
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        # Remove contiguous duplicates (since merged cells in docx return the same text for each cell in the merge)
        cleaned_cells = []
        for c in cells:
            if not cleaned_cells or cleaned_cells[-1] != c:
                cleaned_cells.append(c)
        print(f"  Row {r_idx}: {cleaned_cells}")
